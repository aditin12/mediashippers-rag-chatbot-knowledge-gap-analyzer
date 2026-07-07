"""
query_logger.py
Logs unanswered/low-confidence/off-topic queries to PostgreSQL.
Also exports to Excel sheet on demand.
"""
from datetime import datetime, timezone
from pathlib import Path
from database.postgres import SessionLocal
from database.models import UnansweredQuery, QueryLog

KNOWLEDGE_GAP_THRESHOLD = 0.3

EXPORT_DIR = Path(__file__).parent.parent / "exports"

UNCERTAINTY_PHRASES = [
    "i could not find", "could not find this information",
    "i don't know", "not sure", "no information",
    "cannot find", "can't find", "not mentioned", "not available",
    "no details", "unclear", "not covered", "not specified",
    "the context does not", "no relevant information",
    "i'm unable", "not in the documentation", "insufficient",
]

def detect_confidence(answer: str, retrieved_count: int):
    if retrieved_count == 0:
        return "none", "No relevant documents retrieved"
    for p in UNCERTAINTY_PHRASES:
        if p in answer.lower():
            return "low", f"Answer contains: '{p}'"
    if len(answer.strip()) < 80:
        return "low", "Answer too short"
    return "ok", ""

def log_query(question: str, response: str, retrieved_count: int, out_of_scope: bool = False):
    try:
        if out_of_scope:
            confidence, reason = "out_of_scope", "Not related to MediaShippers platform"
        else:
            confidence, reason = detect_confidence(response, retrieved_count)
            if confidence == "ok":
                return

        now = datetime.now(timezone.utc)
        db  = SessionLocal()
        try:
            existing = db.query(UnansweredQuery).filter(
                UnansweredQuery.question == question.strip()
            ).first()

            if existing:
                existing.frequency += 1
                existing.last_seen  = now
                existing.response   = response
                db.commit()
                freq, updated = existing.frequency, True
            else:
                db.add(UnansweredQuery(
                    question=question.strip(), response=response,
                    confidence=confidence, reason=reason,
                    frequency=1, first_seen=now, last_seen=now,
                ))
                db.commit()
                freq, updated = 1, False

            tag  = "🔄 UPDATED" if updated else "🚨 NEW    "
            conf = {"low":"⚠️  LOW","none":"❌  NONE","out_of_scope":"🚫  OFF-TOPIC"}.get(confidence, confidence)
            print(f"\n{'─'*55}")
            print(f"  {tag}  {conf}  (asked {freq}x)  [PostgreSQL]")
            print(f"  ❓ {question[:80]}")
            print(f"  📋 {reason}")
            print(f"{'─'*55}\n")
        finally:
            db.close()
    except Exception as e:
        print(f"⚠️  Logger error: {e}")

def get_all_unanswered():
    db = SessionLocal()
    try:
        return db.query(UnansweredQuery).order_by(
            UnansweredQuery.frequency.desc(),
            UnansweredQuery.last_seen.desc()
        ).all()
    finally:
        db.close()

def print_report():
    rows = get_all_unanswered()
    if not rows:
        print("\n📊  No unanswered queries logged yet.\n")
        return
    print("\n" + "═"*100)
    print(f"  📊  UNANSWERED QUERIES — {len(rows)} unique  [PostgreSQL]")
    print("═"*100)
    print(f"  {'#':<4} {'TYPE':<18} {'TIMES':<7} {'FIRST SEEN':<22} {'LAST SEEN':<22} QUESTION")
    print("─"*100)
    for i, r in enumerate(rows, 1):
        conf = {"low":"⚠️  LOW","none":"❌  NONE","out_of_scope":"🚫  OFF-TOPIC"}.get(r.confidence, r.confidence)
        q  = r.question[:35] + "..." if len(r.question) > 35 else r.question
        fs = str(r.first_seen)[:19] if r.first_seen else "—"
        ls = str(r.last_seen)[:19]  if r.last_seen  else "—"
        print(f"  {i:<4} {conf:<18} {r.frequency:<7} {fs:<22} {ls:<22} {q}")
    print("═"*100 + "\n")


# ── Knowledge Gap Analytics ─────────────────────────────────────────────────
# Logs EVERY chat question (relevant or not, answered or not) so the
# dashboard can compute total/relevant/irrelevant/answered/gap counts.

def classify_query(score: float, answer: str, retrieved_count: int, out_of_scope: bool):
    """Decide which analytics bucket a query falls into."""
    if out_of_scope or score < KNOWLEDGE_GAP_THRESHOLD:
        return "irrelevant"
    confidence, _ = detect_confidence(answer, retrieved_count)
    if confidence == "ok":
        return "answered_relevant"
    return "knowledge_gap"


def log_analytics(question: str, answer: str, score: float,
                   retrieved_count: int = 0, out_of_scope: bool = False):
    """Insert/update a row in query_logs for the analytics dashboard."""
    try:
        category = classify_query(score, answer, retrieved_count, out_of_scope)
        now = datetime.now(timezone.utc)
        db = SessionLocal()
        try:
            existing = db.query(QueryLog).filter(
                QueryLog.question == question.strip()
            ).first()
            if existing:
                existing.frequency += 1
                existing.last_seen  = now
                existing.answer     = answer
                existing.score      = score
                existing.category   = category
                db.commit()
            else:
                db.add(QueryLog(
                    question=question.strip(), answer=answer, score=score,
                    category=category, frequency=1, first_seen=now, last_seen=now,
                ))
                db.commit()
        finally:
            db.close()
    except Exception as e:
        import traceback
        print(f"⚠️  Analytics logger error: {e}")
        traceback.print_exc()


def get_analytics_summary(top_n: int = 5, recent_n: int = 5):
    """Compute the Knowledge Gap Analytics dashboard data. Never raises."""
    empty = {
        "total_questions": 0, "relevant_questions": 0, "irrelevant_questions": 0,
        "answered_relevant": 0, "knowledge_gaps": 0, "answer_rate": 0.0,
        "top_unanswered": [], "recent_unanswered": [],
    }
    try:
        db = SessionLocal()
        try:
            rows = db.query(QueryLog).all()
            total      = sum(r.frequency for r in rows)
            relevant   = sum(r.frequency for r in rows if r.category in ("answered_relevant", "knowledge_gap"))
            irrelevant = sum(r.frequency for r in rows if r.category == "irrelevant")
            answered   = sum(r.frequency for r in rows if r.category == "answered_relevant")
            gaps       = [r for r in rows if r.category == "knowledge_gap"]
            gap_count  = sum(r.frequency for r in gaps)
            answer_rate = round((answered / total * 100), 1) if total else 0.0
            top_unanswered    = sorted(gaps, key=lambda r: r.frequency, reverse=True)[:top_n]
            recent_unanswered = sorted(gaps, key=lambda r: r.last_seen or r.first_seen, reverse=True)[:recent_n]
            return {
                "total_questions": total, "relevant_questions": relevant,
                "irrelevant_questions": irrelevant, "answered_relevant": answered,
                "knowledge_gaps": gap_count, "answer_rate": answer_rate,
                "top_unanswered":    [{"question": r.question, "frequency": r.frequency} for r in top_unanswered],
                "recent_unanswered": [{"question": r.question, "frequency": r.frequency,
                                       "last_seen": str(r.last_seen)} for r in recent_unanswered],
            }
        finally:
            db.close()
    except Exception as e:
        print(f"⚠️  get_analytics_summary error: {e}")
        return empty


def get_activity_data(from_date: str = None, to_date: str = None):
    """Return daily and weekly activity data for charts."""
    from collections import defaultdict
    from datetime import datetime, timedelta

    empty = {"daily": {"labels": [], "questions": [], "answered": [], "unanswered": []},
             "weekly": {"labels": [], "questions": [], "answered": [], "unanswered": []}}

    try:
        db = SessionLocal()
        try:
            rows = db.query(QueryLog).all()
            if not rows:
                return empty

            # Parse filter dates (plain date comparison, no timezone math)
            fd = datetime.strptime(from_date, "%Y-%m-%d").date() if from_date else None
            td = datetime.strptime(to_date,   "%Y-%m-%d").date() if to_date   else None

            filtered = []
            for r in rows:
                ts = r.last_seen or r.first_seen
                if not ts:
                    continue
                # Strip timezone so .date() always works
                if hasattr(ts, "tzinfo") and ts.tzinfo is not None:
                    ts = ts.replace(tzinfo=None)
                d = ts.date()
                if fd and d < fd:
                    continue
                if td and d > td:
                    continue
                filtered.append((r, ts))

            if not filtered:
                return empty

            # Daily aggregation
            daily = defaultdict(lambda: {"questions": 0, "answered": 0, "unanswered": 0})
            for r, ts in filtered:
                day = ts.strftime("%d %b")
                daily[day]["questions"]  += r.frequency
                if r.category == "answered_relevant":
                    daily[day]["answered"]  += r.frequency
                else:
                    daily[day]["unanswered"] += r.frequency

            sorted_days = sorted(daily.keys(),
                                 key=lambda d: datetime.strptime(d + " 2026", "%d %b %Y"))
            daily_data = {
                "labels":     sorted_days,
                "questions":  [daily[d]["questions"]  for d in sorted_days],
                "answered":   [daily[d]["answered"]   for d in sorted_days],
                "unanswered": [daily[d]["unanswered"] for d in sorted_days],
            }

            # Weekly aggregation
            weekly = defaultdict(lambda: {"questions": 0, "answered": 0, "unanswered": 0})
            for r, ts in filtered:
                week_start = ts.date() - timedelta(days=ts.weekday())
                week_end   = week_start + timedelta(days=6)
                label = f"{week_start.strftime('%d %b')} - {week_end.strftime('%d %b')}"
                weekly[label]["questions"]  += r.frequency
                if r.category == "answered_relevant":
                    weekly[label]["answered"]  += r.frequency
                else:
                    weekly[label]["unanswered"] += r.frequency

            sorted_weeks = sorted(weekly.keys(),
                                  key=lambda w: datetime.strptime(
                                      w.split(" - ")[0] + " 2026", "%d %b %Y"))
            weekly_data = {
                "labels":     sorted_weeks,
                "questions":  [weekly[w]["questions"]  for w in sorted_weeks],
                "answered":   [weekly[w]["answered"]   for w in sorted_weeks],
                "unanswered": [weekly[w]["unanswered"] for w in sorted_weeks],
            }

            return {"daily": daily_data, "weekly": weekly_data}
        finally:
            db.close()
    except Exception as e:
        import traceback
        traceback.print_exc()
        return empty


    """Compute the Knowledge Gap Analytics dashboard data. Never raises —
    always returns a valid shape so the frontend never shows 'undefined'."""
    empty = {
        "total_questions": 0, "relevant_questions": 0, "irrelevant_questions": 0,
        "answered_relevant": 0, "knowledge_gaps": 0, "answer_rate": 0.0,
        "top_unanswered": [], "recent_unanswered": [],
    }
    try:
        db = SessionLocal()
        try:
            rows = db.query(QueryLog).all()

            total      = sum(r.frequency for r in rows)
            relevant   = sum(r.frequency for r in rows if r.category in ("answered_relevant", "knowledge_gap"))
            irrelevant = sum(r.frequency for r in rows if r.category == "irrelevant")
            answered   = sum(r.frequency for r in rows if r.category == "answered_relevant")
            gaps       = [r for r in rows if r.category == "knowledge_gap"]
            gap_count  = sum(r.frequency for r in gaps)
            answer_rate = round((answered / total * 100), 1) if total else 0.0

            top_unanswered = sorted(gaps, key=lambda r: r.frequency, reverse=True)[:top_n]
            recent_unanswered = sorted(gaps, key=lambda r: r.last_seen or r.first_seen, reverse=True)[:recent_n]

            return {
                "total_questions":     total,
                "relevant_questions":  relevant,
                "irrelevant_questions": irrelevant,
                "answered_relevant":   answered,
                "knowledge_gaps":      gap_count,
                "answer_rate":         answer_rate,
                "top_unanswered": [
                    {"question": r.question, "frequency": r.frequency} for r in top_unanswered
                ],
                "recent_unanswered": [
                    {"question": r.question, "frequency": r.frequency,
                     "last_seen": str(r.last_seen)} for r in recent_unanswered
                ],
            }
        finally:
            db.close()
    except Exception as e:
        print(f"⚠️  get_analytics_summary error: {e}")
        return empty


def export_to_excel() -> str:
    """Export unanswered queries to Excel. Returns path to saved file."""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

        rows = get_all_unanswered()
        EXPORT_DIR.mkdir(exist_ok=True)
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        path      = EXPORT_DIR / f"unanswered_queries_{timestamp}.xlsx"

        wb = openpyxl.Workbook()

        # ── Sheet 1: All unanswered queries ──────────────────────────────────
        ws1 = wb.active
        ws1.title = "Unanswered Queries"

        headers = ["#", "Type", "Question", "Reason", "Times Asked", "First Seen", "Last Seen"]
        header_fill   = PatternFill("solid", fgColor="1E3A5F")
        header_font   = Font(bold=True, color="FFFFFF", size=11)
        border_style  = Border(
            left=Side(style="thin"), right=Side(style="thin"),
            top=Side(style="thin"), bottom=Side(style="thin")
        )
        center = Alignment(horizontal="center", vertical="center")

        # Write headers
        for col, h in enumerate(headers, 1):
            cell = ws1.cell(row=1, column=col, value=h)
            cell.font      = header_font
            cell.fill      = header_fill
            cell.alignment = center
            cell.border    = border_style

        # Row fill colors
        fill_oos  = PatternFill("solid", fgColor="FFE0E0")
        fill_low  = PatternFill("solid", fgColor="FFF3CD")
        fill_none = PatternFill("solid", fgColor="F8D7DA")
        fill_alt  = PatternFill("solid", fgColor="F0F4FF")

        for i, r in enumerate(rows, 1):
            conf_label = {
                "low":          "⚠️ LOW CONFIDENCE",
                "none":         "❌ NO ANSWER",
                "out_of_scope": "🚫 OFF-TOPIC",
            }.get(r.confidence, r.confidence)

            row_fill = {"low": fill_low, "none": fill_none, "out_of_scope": fill_oos}.get(
                r.confidence, fill_alt if i % 2 == 0 else None
            )

            row_data = [
                i, conf_label, r.question, r.reason or "", r.frequency,
                str(r.first_seen)[:19] if r.first_seen else "",
                str(r.last_seen)[:19]  if r.last_seen  else "",
            ]

            for col, val in enumerate(row_data, 1):
                cell = ws1.cell(row=i+1, column=col, value=val)
                cell.border = border_style
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                if col in (1, 5): cell.alignment = Alignment(horizontal="center", vertical="center")
                if row_fill: cell.fill = row_fill

        # Column widths
        ws1.column_dimensions["A"].width = 5
        ws1.column_dimensions["B"].width = 22
        ws1.column_dimensions["C"].width = 50
        ws1.column_dimensions["D"].width = 35
        ws1.column_dimensions["E"].width = 12
        ws1.column_dimensions["F"].width = 20
        ws1.column_dimensions["G"].width = 20
        ws1.row_dimensions[1].height = 20

        # ── Sheet 2: Summary by type ─────────────────────────────────────────
        ws2       = wb.create_sheet("Summary")
        ws2["A1"] = "MediaShippers Chatbot — Unanswered Query Summary"
        ws2["A1"].font = Font(bold=True, size=14, color="1E3A5F")
        ws2["A2"] = f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
        ws2["A2"].font = Font(italic=True, color="666666")

        ws2["A4"] = "Category"
        ws2["B4"] = "Count"
        ws2["C4"] = "% of Total"
        for col in ["A4","B4","C4"]:
            ws2[col].font = Font(bold=True, color="FFFFFF")
            ws2[col].fill = PatternFill("solid", fgColor="1E3A5F")
            ws2[col].alignment = center

        counts = {"out_of_scope": 0, "low": 0, "none": 0}
        for r in rows:
            if r.confidence in counts:
                counts[r.confidence] += 1

        total = len(rows)
        labels = {"out_of_scope": "🚫 Off-Topic", "low": "⚠️ Low Confidence", "none": "❌ No Answer"}
        fills  = {"out_of_scope": fill_oos, "low": fill_low, "none": fill_none}

        for row_i, (key, label) in enumerate(labels.items(), 5):
            count = counts[key]
            pct   = f"{(count/total*100):.1f}%" if total > 0 else "0%"
            ws2.cell(row=row_i, column=1, value=label).fill = fills[key]
            ws2.cell(row=row_i, column=2, value=count).fill = fills[key]
            ws2.cell(row=row_i, column=3, value=pct).fill   = fills[key]
            for col in range(1, 4):
                ws2.cell(row=row_i, column=col).border = border_style
                ws2.cell(row=row_i, column=col).alignment = center

        ws2.cell(row=8, column=1, value="TOTAL").font = Font(bold=True)
        ws2.cell(row=8, column=2, value=total).font   = Font(bold=True)
        ws2.column_dimensions["A"].width = 25
        ws2.column_dimensions["B"].width = 12
        ws2.column_dimensions["C"].width = 14

        # Most frequent questions
        ws2["A10"] = "Top 5 Most Asked Unanswered Questions"
        ws2["A10"].font = Font(bold=True, size=12, color="1E3A5F")
        ws2["A11"] = "Rank"
        ws2["B11"] = "Question"
        ws2["C11"] = "Times Asked"
        for col in ["A11","B11","C11"]:
            ws2[col].font = Font(bold=True, color="FFFFFF")
            ws2[col].fill = PatternFill("solid", fgColor="1E3A5F")
            ws2[col].alignment = center

        top5 = sorted(rows, key=lambda x: x.frequency, reverse=True)[:5]
        for row_i, r in enumerate(top5, 12):
            ws2.cell(row=row_i, column=1, value=row_i-11).alignment = center
            ws2.cell(row=row_i, column=2, value=r.question[:80])
            ws2.cell(row=row_i, column=3, value=r.frequency).alignment = center
            for col in range(1, 4):
                ws2.cell(row=row_i, column=col).border = border_style
        ws2.column_dimensions["B"].width = 60

        wb.save(path)
        print(f"\n📊 Excel exported → {path}\n")
        return str(path)

    except ImportError:
        print("⚠️  openpyxl not installed. Run: pip install openpyxl")
        return ""
    except Exception as e:
        print(f"⚠️  Excel export error: {e}")
        return ""
